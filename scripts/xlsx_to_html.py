"""Render a reconciliation .xlsx (Summary + Consolidated Data) to a single
self-contained HTML file for viewing without a spreadsheet app.

    python scripts/xlsx_to_html.py out/report_after_fix.xlsx out/report_after_fix.html
"""
import sys
import html
import openpyxl


def render_sheet(ws, cap=None):
    rows_html = []
    n = 0
    for r in ws.iter_rows(values_only=True):
        n += 1
        if cap and n > cap + 1:  # +1 for header
            rows_html.append(
                f'<tr><td colspan="99" style="background:#fffae6">'
                f"… {ws.max_row - cap - 1:,} more rows — see the .xlsx</td></tr>"
            )
            break
        if all(c is None or c == "" for c in r):
            rows_html.append('<tr class="blank"><td colspan="99"></td></tr>')
            continue
        cells = []
        for c in r:
            if c is None:
                cells.append("<td></td>")
                continue
            if isinstance(c, (int, float)):
                cls = "num neg" if c < 0 else "num"
                cells.append(f'<td class="{cls}">{c:,.2f}</td>')
            else:
                cells.append(f"<td>{html.escape(str(c))}</td>")
        rows_html.append("<tr>" + "".join(cells) + "</tr>")
    return "\n".join(rows_html)


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=False, data_only=True)
    parts = [
        "<!doctype html><meta charset=utf-8>",
        f"<title>{html.escape(src)}</title>",
        """<style>
        body{font:13px/1.4 system-ui,Segoe UI,Arial;margin:24px;color:#1a1a1a}
        h2{margin:28px 0 8px}
        table{border-collapse:collapse;margin-bottom:24px}
        td{border:1px solid #d0d0d0;padding:3px 8px;white-space:nowrap}
        tr:first-child td{background:#f0f0f0;font-weight:600}
        .num{text-align:right;font-variant-numeric:tabular-nums}
        .neg{color:#b00020}
        .blank td{border:0;height:8px}
        .wrap{overflow:auto;max-width:100%;border:1px solid #eee}
        </style>""",
    ]
    cap = int(sys.argv[3]) if len(sys.argv) > 3 else 500
    for name in wb.sheetnames:
        parts.append(f"<h2>{html.escape(name)}</h2>")
        parts.append('<div class="wrap"><table>')
        parts.append(render_sheet(wb[name], cap=None if name == "Summary" else cap))
        parts.append("</table></div>")
    open(dst, "w", encoding="utf-8").write("\n".join(parts))
    print("wrote", dst)


if __name__ == "__main__":
    main()
